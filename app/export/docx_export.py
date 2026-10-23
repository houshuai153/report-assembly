"""Word (.docx) 导出：python-docx。"""

from __future__ import annotations

import io

from docx import Document


def _add_kv_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for k, v in rows:
        cells = table.add_row().cells
        cells[0].text = str(k)
        cells[1].text = str(v)


def build_docx(report: dict) -> bytes:
    meta = report["metadata"]
    doc = Document()

    doc.add_heading(f"{meta['report_type']} — 报告初稿", level=0)
    doc.add_paragraph(f"Study ID: {report['study_id']}　　状态: {report['state']}")

    doc.add_heading("研究信息", level=1)
    _add_kv_table(doc, [
        ("研究编号", meta["study_id"]),
        ("分析物", meta["analyte"]),
        ("基质", meta["matrix"]),
        ("物种", meta["species"]),
        ("方法平台", meta["method_platform"]),
        ("报告类型", meta["report_type"]),
    ])

    doc.add_heading("PK 参数（确定性计算）", level=1)
    headers = ["受试者", "Cmax", "Tmax", "AUC_last", "AUC_inf", "T1/2"]
    pt = doc.add_table(rows=1, cols=len(headers))
    pt.style = "Table Grid"
    for i, h in enumerate(headers):
        pt.rows[0].cells[i].text = h
    for subj, p in report["pk_by_subject"].items():
        cells = pt.add_row().cells
        cells[0].text = subj
        cells[1].text = str(p.get("Cmax"))
        cells[2].text = str(p.get("Tmax"))
        cells[3].text = str(p.get("AUC_last"))
        cells[4].text = "" if p.get("AUC_inf") is None else str(p["AUC_inf"])
        cells[5].text = "" if p.get("T1_2") is None else str(p["T1_2"])

    doc.add_heading("质控统计（确定性计算）", level=1)
    qc = report["qc_summary"]
    _add_kv_table(doc, [
        ("总样本数", qc.get("total_samples")),
        ("低于 LLOQ (BLQ)", qc.get("below_loq")),
        ("可定量样本", qc.get("measurable")),
        ("分析批次", qc.get("runs")),
        ("受试者数", qc.get("subjects")),
        ("LLOQ (ng/mL)", qc.get("lloq")),
    ])

    doc.add_heading("叙述章节（AI 受控生成）", level=1)
    for key, sec in report["sections"].items():
        doc.add_heading(sec["title"], level=2)
        doc.add_paragraph(sec["text"])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
